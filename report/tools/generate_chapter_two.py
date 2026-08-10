from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "report" / "Chapter_Two_Source.md"
OUTPUT = ROOT / "report" / "Chapter_Two_Phishing_Digital_Medium_Detection_System.docx"
SVG_SOURCE = ROOT / "report" / "diagrams" / "PhishGuard_AI_Conceptual_Framework.svg"


def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_margins(cell, top=65, start=65, bottom=65, end=65):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_border(cell, color="7A8FA3", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def allow_row_to_break(row, allow=True):
    tr_pr = row._tr.get_or_add_trPr()
    existing = tr_pr.find(qn("w:cantSplit"))
    if allow and existing is not None:
        tr_pr.remove(existing)
    elif not allow and existing is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_fixed_cell_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_layout_fixed(table):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=12)


def set_page_number_start(section, value=1):
    pg_num_type = section._sectPr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        section._sectPr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(value))


def set_standard_page(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.45)


def set_landscape_page(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.3)


def set_section_footer_link(section):
    section.footer.is_linked_to_previous = True
    section.header.is_linked_to_previous = True
    # python-docx clones the prior section properties when a new section is
    # added. Remove the cloned start value so numbering continues instead of
    # restarting at page 1 after landscape/portrait transitions.
    pg_num_type = section._sectPr.find(qn("w:pgNumType"))
    if pg_num_type is not None:
        section._sectPr.remove(pg_num_type)


def add_inline_markup(paragraph, text, size=12, color=None):
    """Translate the limited **bold** and *italic* syntax used by the source."""
    token_re = re.compile(r"(\*\*.*?\*\*|(?<!\*)\*[^*]+?\*|`[^`]+?`)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color=color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True, color=color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, color=color)


def configure_styles(document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    style_specs = {
        "PCU Section Heading": (12, True, WD_ALIGN_PARAGRAPH.LEFT, 12),
        "PCU Subsection Heading": (12, True, WD_ALIGN_PARAGRAPH.LEFT, 8),
        "PCU Caption": (10, False, WD_ALIGN_PARAGRAPH.CENTER, 6),
        "PCU Diagram": (9, False, WD_ALIGN_PARAGRAPH.CENTER, 0),
    }
    for name, (size, bold, align, before) in style_specs.items():
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.paragraph_format.alignment = align
        style.paragraph_format.line_spacing = 1.0 if name in ("PCU Caption", "PCU Diagram") else 1.5
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = name != "PCU Diagram"

    if "PCU Reference" not in styles:
        style = styles.add_style("PCU Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["PCU Reference"]
    style.font.name = "Times New Roman"
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    style.paragraph_format.left_indent = Inches(0.5)
    style.paragraph_format.first_line_indent = Inches(-0.5)
    style.paragraph_format.space_after = Pt(0)


def add_cover(document):
    section = document.sections[0]
    set_standard_page(section)
    v_align = OxmlElement("w:vAlign")
    v_align.set(qn("w:val"), "center")
    section._sectPr.append(v_align)

    entries = [
        ("PHISHING DIGITAL MEDIUM DETECTION SYSTEM USING MACHINE LEARNING", 14, True, 24),
        ("CHAPTER TWO: LITERATURE REVIEW", 13, True, 18),
        ("BY", 12, True, 16),
        ("HONOUR JESUTOFAYE JESUTOFAYE", 13, True, 4),
        ("MATRICULATION NUMBER: [INSERT MATRICULATION NUMBER]", 12, True, 24),
        (
            "A PROJECT REPORT SUBMITTED TO THE DEPARTMENT OF COMPUTER SCIENCE, "
            "PRECIOUS CORNERSTONE UNIVERSITY, IBADAN, IN PARTIAL FULFILMENT OF "
            "THE REQUIREMENTS FOR THE AWARD OF THE DEGREE OF BACHELOR OF SCIENCE "
            "(B.Sc.) IN COMPUTER SCIENCE",
            12,
            True,
            22,
        ),
        ("SUPERVISOR: DR. OSUNTOKUN", 12, True, 20),
        ("[MONTH, YEAR]", 12, True, 0),
    ]
    for text, size, bold, after in entries:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(after)
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold)


def begin_main_section(document):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    set_standard_page(section)
    set_page_number_start(section, 1)
    section.footer.is_linked_to_previous = False
    section.header.is_linked_to_previous = False
    add_page_number(section.footer.paragraphs[0])
    return section


def add_chapter_title(document, text, size=14, space_after=0):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=True)
    return paragraph


def add_heading(document, text, subsection=False):
    style = "PCU Subsection Heading" if subsection else "PCU Section Heading"
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.page_break_before = False
    run = paragraph.add_run(text)
    set_run_font(run, size=12, bold=True)
    return paragraph


def add_body_paragraph(document, text):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Inches(0.5)
    add_inline_markup(paragraph, text)
    return paragraph


def add_bullet_item(document, text):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    run = paragraph.add_run("•\t")
    set_run_font(run)
    add_inline_markup(paragraph, text)
    return paragraph


def add_numbered_item(document, text):
    match = re.match(r"^(\d+)\.\s+(.*)$", text)
    if not match:
        return add_body_paragraph(document, text)
    number, body = match.groups()
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.3)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))
    run = paragraph.add_run(f"{number}.\t")
    set_run_font(run)
    add_inline_markup(paragraph, body)
    return paragraph


def add_reference(document, text):
    paragraph = document.add_paragraph(style="PCU Reference")
    add_inline_markup(paragraph, text)
    return paragraph


def add_caption(document, label, text, before=False):
    paragraph = document.add_paragraph(style="PCU Caption")
    paragraph.paragraph_format.space_before = Pt(0 if before else 5)
    paragraph.paragraph_format.space_after = Pt(5 if before else 0)
    run = paragraph.add_run(f"{label}: ")
    set_run_font(run, size=10, bold=True)
    add_inline_markup(paragraph, text, size=10)
    return paragraph


def parse_markdown_table(table_lines):
    rows = []
    for line in table_lines:
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        parts = [part.strip() for part in stripped.strip("|").split("|")]
        if parts:
            rows.append(parts)
    if not rows:
        raise ValueError("No rows found for Table 2.1")
    width = len(rows[0])
    if any(len(row) != width for row in rows):
        raise ValueError("Inconsistent Table 2.1 column count")
    return rows


def add_comparative_table(document, rows):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    set_landscape_page(section)
    set_section_footer_link(section)

    add_caption(
        document,
        "Table 2.1",
        "Comparative analysis of selected phishing email and SMS studies (2024–2026)",
        before=True,
    )

    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_layout_fixed(table)
    widths = [1.22, 1.05, 1.55, 1.62, 1.68, 1.55, 1.53]

    for row_index, (word_row, docx_row) in enumerate(zip(rows, table.rows)):
        allow_row_to_break(docx_row, allow=True)
        if row_index == 0:
            set_repeat_table_header(docx_row)
        for col_index, (text, cell) in enumerate(zip(word_row, docx_row.cells)):
            set_fixed_cell_width(cell, widths[col_index])
            set_cell_margins(cell, top=45, start=45, bottom=45, end=45)
            set_cell_border(cell, color="6F8295", size="6")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shade_cell(cell, "244A68")
            elif row_index % 2 == 0:
                shade_cell(cell, "EAF0F5")
            else:
                shade_cell(cell, "FFFFFF")

            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(text)
            set_run_font(
                run,
                size=7.2 if row_index else 7.5,
                bold=row_index == 0,
                color=(255, 255, 255) if row_index == 0 else None,
            )

    note = document.add_paragraph(style="PCU Caption")
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(0)
    run = note.add_run("Note. ")
    set_run_font(run, size=8.5, italic=True)
    add_inline_markup(
        note,
        "Metrics are reproduced from the cited studies and are not directly comparable across different datasets or evaluation designs.",
        size=8.5,
    )

    next_section = document.add_section(WD_SECTION.NEW_PAGE)
    set_standard_page(next_section)
    set_section_footer_link(next_section)


def set_diagram_cell(cell, fill, title, body="", title_color=(16, 42, 67), body_color=(36, 59, 83)):
    shade_cell(cell, fill)
    set_cell_border(cell, color="74889A", size="7")
    set_cell_margins(cell, top=75, start=70, bottom=75, end=70)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)
    title_run = paragraph.add_run(title)
    set_run_font(title_run, size=9, bold=True, color=title_color)
    if body:
        body_run = paragraph.add_run(f"\n{body}")
        set_run_font(body_run, size=7.7, color=body_color)


def merge_row(row):
    merged = row.cells[0]
    for cell in row.cells[1:]:
        merged = merged.merge(cell)
    return row.cells[0]


def set_stage_row(row, text):
    cell = merge_row(row)
    shade_cell(cell, "D9E5EF")
    set_cell_border(cell, color="315A7D", size="8")
    set_cell_margins(cell, top=40, start=55, bottom=40, end=55)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, size=8.5, bold=True, color=(49, 90, 125))


def set_arrow_row(row):
    cell = merge_row(row)
    set_cell_border(cell, color="FFFFFF", size="0")
    set_cell_margins(cell, top=5, start=0, bottom=5, end=0)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run("↓")
    set_run_font(run, size=13, bold=True, color=(49, 90, 125))


def add_conceptual_framework(document):
    table = document.add_table(rows=17, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_layout_fixed(table)
    for row in table.rows:
        allow_row_to_break(row, allow=False)
        for cell in row.cells:
            set_fixed_cell_width(cell, 1.66)

    set_stage_row(table.rows[0], "1  USER INPUT LAYER")
    email_cell = table.rows[1].cells[0].merge(table.rows[1].cells[1])
    sms_cell = table.rows[1].cells[2].merge(table.rows[1].cells[3])
    set_diagram_cell(email_cell, "E7F0FA", "EMAIL MODE", "Sender email • Optional subject • Message body")
    set_diagram_cell(sms_cell, "E8F6F3", "SMS MODE", "Phone / short code / sender ID • Message body")
    set_arrow_row(table.rows[2])

    set_stage_row(table.rows[3], "2  VALIDATION AND CHANNEL ROUTING")
    validation_cell = merge_row(table.rows[4])
    set_diagram_cell(
        validation_cell,
        "315A7D",
        "VALIDATE REQUIRED FIELDS AND SELECT THE MATCHING SAVED MODEL",
        "Preserve email/SMS field differences and reject invalid submissions",
        title_color=(255, 255, 255),
        body_color=(255, 255, 255),
    )
    set_arrow_row(table.rows[5])

    set_stage_row(table.rows[6], "3  PARALLEL EVIDENCE LAYER")
    evidence = [
        ("EEF4FB", "ML TEXT EVIDENCE", "TF-IDF + Logistic Regression\nBounded contribution"),
        ("F4F0FB", "CONTENT RULES", "Urgency • threats • credentials\nOTP • payment • reward"),
        ("EAF7F4", "SENDER / DOMAIN", "Format • brand mismatch\nLook-alike heuristics"),
        ("FFF4E5", "URL ANALYSIS", "Shorteners • raw IP • @\nPunycode • selected TLDs"),
    ]
    for cell, (fill, title, body) in zip(table.rows[7].cells, evidence):
        set_diagram_cell(cell, fill, title, body)
    set_arrow_row(table.rows[8])

    set_stage_row(table.rows[9], "4  HYBRID RISK AGGREGATION")
    aggregate_cell = merge_row(table.rows[10])
    set_diagram_cell(
        aggregate_cell,
        "243B53",
        "COMBINE EVIDENCE INTO A BOUNDED 0–100 SCORE",
        "ML contribution + weighted findings + contextual safeguards; triage score, not a calibrated probability",
        title_color=(255, 255, 255),
        body_color=(255, 255, 255),
    )
    set_arrow_row(table.rows[11])

    set_stage_row(table.rows[12], "5  RISK DECISION")
    low_cell = table.rows[13].cells[0]
    suspicious_cell = table.rows[13].cells[1].merge(table.rows[13].cells[2])
    high_cell = table.rows[13].cells[3]
    set_diagram_cell(low_cell, "E6F4EA", "LOW RISK", "0–39")
    set_diagram_cell(suspicious_cell, "FFF6D6", "SUSPICIOUS", "40–69")
    set_diagram_cell(high_cell, "FDE9E7", "HIGH RISK", "70–100")
    set_arrow_row(table.rows[14])

    set_stage_row(table.rows[15], "6  DECISION-SUPPORT OUTPUT")
    output = [
        ("VERDICT + SCORE", "Risk-oriented result"),
        ("EVIDENCE", "Findings + safe signals"),
        ("USER GUIDANCE", "Explanations + actions"),
        ("PERSISTENCE", "Local scan record"),
    ]
    for cell, (title, body) in zip(table.rows[16].cells, output):
        set_diagram_cell(cell, "EDF2F7", title, body)

    evaluation = document.add_table(rows=1, cols=1)
    evaluation.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_layout_fixed(evaluation)
    cell = evaluation.cell(0, 0)
    set_fixed_cell_width(cell, 6.64)
    set_diagram_cell(
        cell,
        "FFF8EA",
        "EVALUATION BOUNDARY (NOT AN AUTOMATIC SELF-LEARNING LOOP)",
        "Verified metrics • cross-dataset • temporal • multilingual • adversarial • user testing",
        title_color=(124, 62, 0),
        body_color=(92, 59, 22),
    )


def add_figure_caption_from_source(document, line):
    match = re.match(r"^\*\*Figure\s+2\.1:\*\*\s*(.*)$", line)
    if not match:
        return False
    add_caption(document, "Figure 2.1", match.group(1), before=False)
    return True


def parse_source(document, lines):
    in_references = False
    skipped_title = False
    index = 0

    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line.startswith("# ") and not skipped_title:
            skipped_title = True
            continue
        if line == "## CHAPTER TWO":
            add_chapter_title(document, "CHAPTER TWO", size=14, space_after=0)
            continue
        if line == "## LITERATURE REVIEW":
            add_chapter_title(document, "LITERATURE REVIEW", size=14, space_after=12)
            continue
        if line == "## REFERENCES":
            document.add_page_break()
            add_chapter_title(document, "REFERENCES", size=14, space_after=12)
            in_references = True
            continue
        if line == "[[TABLE 2.1 START]]":
            table_lines = []
            while index < len(lines) and lines[index].strip() != "[[TABLE 2.1 END]]":
                table_lines.append(lines[index])
                index += 1
            if index >= len(lines):
                raise ValueError("Table 2.1 end marker not found")
            index += 1
            add_comparative_table(document, parse_markdown_table(table_lines))
            continue
        if line == "[[FIGURE 2.1: PHISHGUARD AI CONCEPTUAL FRAMEWORK]]":
            add_conceptual_framework(document)
            continue
        if add_figure_caption_from_source(document, line):
            continue
        if line.startswith("#### "):
            add_heading(document, line[5:], subsection=True)
            continue
        if line.startswith("### "):
            add_heading(document, line[4:], subsection=False)
            continue
        if in_references:
            add_reference(document, line)
            continue
        if line.startswith("- "):
            add_bullet_item(document, line[2:])
            continue
        if re.match(r"^\d+\.\s", line):
            add_numbered_item(document, line)
            continue
        add_body_paragraph(document, line)


def set_document_properties(document):
    props = document.core_properties
    props.title = "Chapter Two — Phishing Digital Medium Detection System Using Machine Learning"
    props.subject = "Final-Year Project Report — Literature Review"
    props.author = "Honour Jesutofaye Jesutofaye"
    props.keywords = (
        "phishing, smishing, email, SMS, machine learning, TF-IDF, Logistic Regression, "
        "hybrid detection, explainable AI, PhishGuard AI"
    )
    props.comments = (
        "Chapter Two review draft based on recent verified literature and the implemented "
        "PhishGuard AI repository. Reported study metrics are not PhishGuard AI results."
    )


def set_update_fields(document):
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not SVG_SOURCE.exists():
        raise FileNotFoundError(SVG_SOURCE)

    document = Document()
    configure_styles(document)
    set_document_properties(document)
    add_cover(document)
    begin_main_section(document)
    parse_source(document, SOURCE.read_text(encoding="utf-8").splitlines())
    set_update_fields(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
