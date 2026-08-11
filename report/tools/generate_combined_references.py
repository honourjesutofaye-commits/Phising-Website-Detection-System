from __future__ import annotations

"""Generate a standalone Word copy of the combined Chapters One–Five references."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

import generate_chapter_five as document_tools

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "report" / "Complete_Report_Source.md"
OUTPUT = ROOT / "report" / "Combined_References_Chapters_One_to_Five.docx"


def main():
    source = SOURCE.read_text(encoding="utf-8")
    marker = "\n## REFERENCES\n"
    if source.count(marker) != 1:
        raise ValueError("Complete report source must contain exactly one reference section")
    references = [entry.strip() for entry in source.split(marker, 1)[1].strip().split("\n\n") if entry.strip()]
    if len(references) != 35:
        raise ValueError(f"Expected 35 consolidated references, found {len(references)}")

    document = Document()
    document_tools.configure_styles(document)
    section = document.sections[0]
    document_tools.set_standard_page(section)
    document_tools.set_page_number_start(section, 1)
    document_tools.add_page_number(section.footer.paragraphs[0])

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(18)
    run = heading.add_run("REFERENCES")
    document_tools.set_run_font(run, size=14, bold=True)

    for reference in references:
        document_tools.add_reference(document, reference)

    properties = document.core_properties
    properties.title = "Combined References — Chapters One to Five"
    properties.subject = "Phishing Digital Medium Detection System Using Machine Learning"
    properties.author = "Honour Jesutofaye Jesutofaye"
    properties.comments = "Alphabetised and deduplicated APA reference list from the complete project report."

    settings = document.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")

    document.save(OUTPUT)
    print(f"{OUTPUT}\n{len(references)} alphabetised, deduplicated references")


if __name__ == "__main__":
    main()
