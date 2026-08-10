from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "report" / "Chapter_One_Source.md"
OUTPUT = ROOT / "report" / "Chapter_One_Phishing_Digital_Medium_Detection_System.docx"


def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=12)


def set_page_number_start(section, value=1):
    pg_num_type = section._sectPr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        section._sectPr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(value))


def add_inline_markup(paragraph, text, size=12):
    """Translate the limited **bold** and *italic* markup used by the source."""
    token_re = re.compile(r"(\*\*.*?\*\*|\*.*?\*)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size)


def set_standard_page(section):
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.45)


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    if "PCU Section Heading" not in styles:
        style = styles.add_style("PCU Section Heading", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["PCU Section Heading"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)
    style.font.bold = True
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.keep_with_next = True

    if "PCU Subsection Heading" not in styles:
        style = styles.add_style("PCU Subsection Heading", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["PCU Subsection Heading"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)
    style.font.bold = True
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.keep_with_next = True

    if "PCU Reference" not in styles:
        style = styles.add_style("PCU Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["PCU Reference"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    style.paragraph_format.left_indent = Inches(0.5)
    style.paragraph_format.first_line_indent = Inches(-0.5)
    style.paragraph_format.space_after = Pt(0)


def add_cover(document):
    section = document.sections[0]
    set_standard_page(section)
    # Vertically centre the title-page block.
    v_align = OxmlElement("w:vAlign")
    v_align.set(qn("w:val"), "center")
    section._sectPr.append(v_align)

    entries = [
        ("PHISHING DIGITAL MEDIUM DETECTION SYSTEM USING MACHINE LEARNING", 14, True, 26),
        ("BY", 12, True, 18),
        ("HONOUR JESUTOFAYE JESUTOFAYE", 13, True, 4),
        ("MATRICULATION NUMBER: [INSERT MATRICULATION NUMBER]", 12, True, 26),
        (
            "A PROJECT REPORT SUBMITTED TO THE DEPARTMENT OF COMPUTER SCIENCE, "
            "PRECIOUS CORNERSTONE UNIVERSITY, IBADAN, IN PARTIAL FULFILMENT OF "
            "THE REQUIREMENTS FOR THE AWARD OF THE DEGREE OF BACHELOR OF SCIENCE "
            "(B.Sc.) IN COMPUTER SCIENCE",
            12,
            True,
            24,
        ),
        ("SUPERVISOR: DR. OSUNTOKUN", 12, True, 22),
        ("[MONTH, YEAR]", 12, True, 0),
    ]

    for text, size, bold, after in entries:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(after)
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold)


def begin_main_section(document):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    set_standard_page(section)
    set_page_number_start(section, 1)
    section.footer.is_linked_to_previous = False
    footer_p = section.footer.paragraphs[0]
    add_page_number(footer_p)
    section.header.is_linked_to_previous = False
    return section


def add_chapter_title(document, text, size=14, space_after=0):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True)
    return p


def add_body_paragraph(document, text):
    p = document.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Inches(0.5)
    add_inline_markup(p, text)
    return p


def add_numbered_item(document, text):
    match = re.match(r"^(\d+)\.\s+(.*)$", text)
    if not match:
        return add_body_paragraph(document, text)
    number, body = match.groups()
    p = document.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))
    r = p.add_run(f"{number}.\t")
    set_run_font(r)
    add_inline_markup(p, body)
    return p


def add_heading(document, text, subsection=False):
    style = "PCU Subsection Heading" if subsection else "PCU Section Heading"
    p = document.add_paragraph(style=style)
    p.paragraph_format.page_break_before = False
    r = p.add_run(text)
    set_run_font(r, size=12, bold=True)
    return p


def add_reference(document, text):
    p = document.add_paragraph(style="PCU Reference")
    add_inline_markup(p, text)
    return p


def parse_source(document, lines):
    in_references = False
    skipped_title = False
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# ") and not skipped_title:
            skipped_title = True
            continue
        if line == "## CHAPTER ONE":
            add_chapter_title(document, "CHAPTER ONE", size=14, space_after=0)
            continue
        if line == "## INTRODUCTION":
            add_chapter_title(document, "INTRODUCTION", size=14, space_after=12)
            continue
        if line == "## REFERENCES":
            document.add_page_break()
            add_chapter_title(document, "REFERENCES", size=14, space_after=12)
            in_references = True
            continue
        if line.startswith("### "):
            add_heading(document, line[4:], subsection=False)
            continue
        if line.startswith("#### "):
            add_heading(document, line[5:], subsection=True)
            continue
        if in_references:
            add_reference(document, line)
            continue
        if re.match(r"^\d+\.\s", line):
            add_numbered_item(document, line)
        else:
            add_body_paragraph(document, line)


def set_document_properties(document):
    props = document.core_properties
    props.title = "Chapter One — Phishing Digital Medium Detection System Using Machine Learning"
    props.subject = "Final-Year Project Report — Chapter One Review Draft"
    props.author = "Honour Jesutofaye Jesutofaye"
    props.keywords = "phishing, smishing, machine learning, TF-IDF, Logistic Regression, Django, PhishGuard AI"
    props.comments = (
        "Chapter One review draft prepared from the implemented PhishGuard AI repository. "
        "Matriculation number and submission date remain editable placeholders."
    )


def set_update_fields(document):
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main():
    document = Document()
    configure_styles(document)
    set_document_properties(document)
    add_cover(document)
    begin_main_section(document)
    parse_source(document, SOURCE.read_text(encoding="utf-8").splitlines())
    set_update_fields(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
