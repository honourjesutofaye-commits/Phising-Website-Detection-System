from __future__ import annotations

from pathlib import Path
import os
import re
import shutil
import tempfile
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
import resvg_py

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "report" / "Chapter_Three_Source.md"
OUTPUT = ROOT / "report" / "Chapter_Three_Phishing_Digital_Medium_Detection_System.docx"
RENDER_DIR = ROOT / ".report-work" / "chapter-three-renders"
DIAGRAM_DIR = ROOT / "report" / "diagrams"

FIGURES = {
    "DEVELOPMENT METHODOLOGY": {
        "number": "3.1",
        "svg": DIAGRAM_DIR / "PhishGuard_AI_Chapter3_Methodology.svg",
        "landscape": True,
        "width": 9.25,
    },
    "USE-CASE MODEL": {
        "number": "3.2",
        "svg": DIAGRAM_DIR / "PhishGuard_AI_Chapter3_Use_Case.svg",
        "landscape": True,
        "width": 7.25,
    },
    "SYSTEM ARCHITECTURE": {
        "number": "3.3",
        "svg": DIAGRAM_DIR / "PhishGuard_AI_Chapter3_Architecture.svg",
        "landscape": True,
        "width": 7.75,
    },
    "DATA-FLOW DIAGRAM": {
        "number": "3.4",
        "svg": DIAGRAM_DIR / "PhishGuard_AI_Chapter3_DFD.svg",
        "landscape": True,
        "width": 9.25,
    },
    "HYBRID ANALYSIS ACTIVITY": {
        "number": "3.5",
        "svg": DIAGRAM_DIR / "PhishGuard_AI_Chapter3_Activity.svg",
        "landscape": False,
        "width": 5.45,
    },
    "DATA MODEL": {
        "number": "3.6",
        "svg": DIAGRAM_DIR / "PhishGuard_AI_Chapter3_Data_Model.svg",
        "landscape": False,
        "width": 5.75,
    },
}

LANDSCAPE_TABLES = {"3.4", "3.5", "3.6"}
TABLE_WIDTHS = {
    "3.1": [1.45, 2.35, 2.45],
    "3.2": [0.65, 3.25, 2.35],
    "3.3": [0.65, 1.25, 4.35],
    "3.4": [1.30, 2.10, 3.20, 3.10],
    "3.5": [1.55, 2.00, 3.00, 3.15],
    "3.6": [1.25, 1.80, 2.70, 3.95],
    "3.7": [1.50, 2.35, 2.40],
    "3.8": [1.95, 1.20, 3.10],
    "3.9": [1.90, 1.45, 2.90],
    "3.10": [1.50, 1.50, 3.25],
    "3.11": [1.25, 2.50, 2.50],
    "3.12": [1.55, 2.05, 2.65],
}


def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None, color=None):
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_margins(cell, top=60, start=60, bottom=60, end=60):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)
    node.set(qn("w:val"), "clear")


def set_cell_border(cell, color="71869A", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:tblHeader"))
    if node is None:
        node = OxmlElement("w:tblHeader")
        tr_pr.append(node)
    node.set(qn("w:val"), "true")


def allow_row_to_break(row, allow=True):
    tr_pr = row._tr.get_or_add_trPr()
    existing = tr_pr.find(qn("w:cantSplit"))
    if allow and existing is not None:
        tr_pr.remove(existing)
    elif not allow and existing is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_fixed_cell_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:tcW"))
    if node is None:
        node = OxmlElement("w:tcW")
        tc_pr.append(node)
    node.set(qn("w:w"), str(int(width_inches * 1440)))
    node.set(qn("w:type"), "dxa")


def set_table_layout_fixed(table):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    node = tbl_pr.find(qn("w:tblLayout"))
    if node is None:
        node = OxmlElement("w:tblLayout")
        tbl_pr.append(node)
    node.set(qn("w:type"), "fixed")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=12)


def set_page_number_start(section, value=1):
    node = section._sectPr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType")
        section._sectPr.append(node)
    node.set(qn("w:start"), str(value))


def set_standard_page(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.45)


def set_landscape_page(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.3)


def continue_numbering(section):
    section.footer.is_linked_to_previous = True
    section.header.is_linked_to_previous = True
    node = section._sectPr.find(qn("w:pgNumType"))
    if node is not None:
        section._sectPr.remove(node)


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    fonts = normal._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    specifications = {
        "PCU Section Heading": (12, True, WD_ALIGN_PARAGRAPH.LEFT, 12, 1.5),
        "PCU Subsection Heading": (12, True, WD_ALIGN_PARAGRAPH.LEFT, 8, 1.5),
        "PCU Caption": (10, False, WD_ALIGN_PARAGRAPH.CENTER, 5, 1.0),
        "PCU Figure": (9, False, WD_ALIGN_PARAGRAPH.CENTER, 0, 1.0),
    }
    for name, (size, bold, alignment, before, spacing) in specifications.items():
        style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.paragraph_format.alignment = alignment
        style.paragraph_format.line_spacing = spacing
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = name != "PCU Figure"

    reference = styles["PCU Reference"] if "PCU Reference" in styles else styles.add_style("PCU Reference", WD_STYLE_TYPE.PARAGRAPH)
    reference.font.name = "Times New Roman"
    reference._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    reference.font.size = Pt(12)
    reference.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    reference.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    reference.paragraph_format.left_indent = Inches(0.5)
    reference.paragraph_format.first_line_indent = Inches(-0.5)
    reference.paragraph_format.space_after = Pt(0)


def add_cover(document):
    section = document.sections[0]
    set_standard_page(section)
    vertical = OxmlElement("w:vAlign")
    vertical.set(qn("w:val"), "center")
    section._sectPr.append(vertical)
    entries = [
        ("PHISHING DIGITAL MEDIUM DETECTION SYSTEM USING MACHINE LEARNING", 14, True, 22),
        ("CHAPTER THREE: RESEARCH METHODOLOGY, SYSTEM ANALYSIS AND DESIGN", 13, True, 16),
        ("BY", 12, True, 14),
        ("HONOUR JESUTOFAYE JESUTOFAYE", 13, True, 4),
        ("MATRICULATION NUMBER: [INSERT MATRICULATION NUMBER]", 12, True, 22),
        (
            "A PROJECT REPORT SUBMITTED TO THE DEPARTMENT OF COMPUTER SCIENCE, "
            "PRECIOUS CORNERSTONE UNIVERSITY, IBADAN, IN PARTIAL FULFILMENT OF "
            "THE REQUIREMENTS FOR THE AWARD OF THE DEGREE OF BACHELOR OF SCIENCE "
            "(B.Sc.) IN COMPUTER SCIENCE",
            12,
            True,
            20,
        ),
        ("SUPERVISOR: DR. OSUNTOKUN", 12, True, 18),
        ("[MONTH, YEAR]", 12, True, 0),
    ]
    for text, size, bold, after in entries:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(after)
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold)


def begin_main_section(document):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    set_standard_page(section)
    set_page_number_start(section, 1)
    section.footer.is_linked_to_previous = False
    section.header.is_linked_to_previous = False
    add_page_number(section.footer.paragraphs[0])


def add_inline_markup(paragraph, text, size=12, color=None):
    text = text.replace("\\(", "").replace("\\)", "")
    token_re = re.compile(r"(\*\*.*?\*\*|(?<!\*)\*[^*]+?\*|`[^`]+?`)")
    position = 0
    for match in token_re.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color=color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True, color=color)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, color=color)


def add_chapter_title(document, text, size=14, space_after=0):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=True)


def add_heading(document, text, subsection=False):
    paragraph = document.add_paragraph(style="PCU Subsection Heading" if subsection else "PCU Section Heading")
    run = paragraph.add_run(text)
    set_run_font(run, size=12, bold=True)


def add_body_paragraph(document, text):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Inches(0.5)
    add_inline_markup(paragraph, text)


def add_bullet_item(document, text):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    run = paragraph.add_run("•\t")
    set_run_font(run)
    add_inline_markup(paragraph, text)


def add_numbered_item(document, text):
    match = re.match(r"^(\d+)\.\s+(.*)$", text)
    if not match:
        add_body_paragraph(document, text)
        return
    number, body = match.groups()
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.3)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))
    run = paragraph.add_run(f"{number}.\t")
    set_run_font(run)
    add_inline_markup(paragraph, body)


def add_reference(document, text):
    paragraph = document.add_paragraph(style="PCU Reference")
    add_inline_markup(paragraph, text)


def add_caption(document, label, text, before=False):
    paragraph = document.add_paragraph(style="PCU Caption")
    paragraph.paragraph_format.space_before = Pt(0 if before else 5)
    paragraph.paragraph_format.space_after = Pt(5 if before else 0)
    paragraph.paragraph_format.keep_with_next = before
    run = paragraph.add_run(f"{label}: ")
    set_run_font(run, size=10, bold=True)
    add_inline_markup(paragraph, text, size=10)


def parse_markdown_table(lines):
    rows = []
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("|"):
            rows.append([part.strip() for part in stripped.strip("|").split("|")])
    if len(rows) < 2:
        raise ValueError("Table contains fewer than two rows")
    width = len(rows[0])
    if any(len(row) != width for row in rows):
        raise ValueError("Inconsistent table column count")
    return rows


def clean_table_text(text):
    return text.replace("`", "").replace("**", "").replace("\\(", "").replace("\\)", "")


def add_table(document, number, caption, rows):
    landscape = number in LANDSCAPE_TABLES
    if landscape:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        set_landscape_page(section)
        continue_numbering(section)

    add_caption(document, f"Table {number}", caption, before=True)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_layout_fixed(table)
    widths = TABLE_WIDTHS[number]
    if len(widths) != len(rows[0]):
        raise ValueError(f"Width specification mismatch for Table {number}")

    font_size = 7.8 if landscape else 8.6
    for row_index, (source_row, output_row) in enumerate(zip(rows, table.rows)):
        allow_row_to_break(output_row, allow=row_index != 0)
        if row_index == 0:
            set_repeat_table_header(output_row)
        for column_index, (value, cell) in enumerate(zip(source_row, output_row.cells)):
            set_fixed_cell_width(cell, widths[column_index])
            set_cell_margins(cell, top=45, start=50, bottom=45, end=50)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            shade_cell(cell, "244A68" if row_index == 0 else ("EAF0F5" if row_index % 2 == 0 else "FFFFFF"))
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_markup(
                paragraph,
                clean_table_text(value),
                size=font_size if row_index else font_size + 0.3,
                color=(255, 255, 255) if row_index == 0 else None,
            )
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True

    if landscape:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        set_standard_page(section)
        continue_numbering(section)


def render_svg(svg_path, png_path):
    png_path.parent.mkdir(parents=True, exist_ok=True)
    png_path.write_bytes(
        resvg_py.svg_to_bytes(
            svg_path=str(svg_path),
            width=1800,
            background="#ffffff",
            font_dirs=["/usr/share/fonts/truetype/dejavu"],
            style_sheet="text { font-family: 'DejaVu Sans' !important; }",
        )
    )


def add_figure(document, key, pending_svg_links):
    figure = FIGURES[key]
    svg_path = figure["svg"]
    png_path = RENDER_DIR / f"figure-{figure['number']}.png"
    render_svg(svg_path, png_path)

    if figure["landscape"]:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        set_landscape_page(section)
        continue_numbering(section)
    else:
        document.add_page_break()

    paragraph = document.add_paragraph(style="PCU Figure")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    shape = run.add_picture(str(png_path), width=Inches(figure["width"]))
    blip = shape._inline.graphic.graphicData.pic.blipFill.blip
    pending_svg_links.append((blip.embed, svg_path))


def finish_figure_section(document, key):
    if FIGURES[key]["landscape"]:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        set_standard_page(section)
        continue_numbering(section)


def parse_source(document, lines, pending_svg_links):
    in_references = False
    skipped_title = False
    pending_figure_key = None
    index = 0

    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line.startswith("# ") and not skipped_title:
            skipped_title = True
            continue
        if line == "## CHAPTER THREE":
            add_chapter_title(document, "CHAPTER THREE", size=14)
            continue
        if line == "## RESEARCH METHODOLOGY, SYSTEM ANALYSIS AND DESIGN":
            add_chapter_title(document, "RESEARCH METHODOLOGY, SYSTEM ANALYSIS AND DESIGN", size=14, space_after=12)
            continue
        if line == "## REFERENCES":
            document.add_page_break()
            add_chapter_title(document, "REFERENCES", size=14, space_after=12)
            in_references = True
            continue

        table_start = re.match(r"^\[\[TABLE\s+(3\.\d+)\s+START\]\]$", line)
        if table_start:
            number = table_start.group(1)
            table_lines = []
            end_marker = f"[[TABLE {number} END]]"
            while index < len(lines) and lines[index].strip() != end_marker:
                table_lines.append(lines[index])
                index += 1
            if index >= len(lines):
                raise ValueError(f"Missing end marker for Table {number}")
            index += 1
            while index < len(lines) and not lines[index].strip():
                index += 1
            if index >= len(lines):
                raise ValueError(f"Missing caption for Table {number}")
            caption_match = re.match(rf"^\*\*Table\s+{re.escape(number)}:\*\*\s*(.*)$", lines[index].strip())
            if not caption_match:
                raise ValueError(f"Expected caption immediately after Table {number}")
            index += 1
            add_table(document, number, caption_match.group(1), parse_markdown_table(table_lines))
            continue

        figure_match = re.match(r"^\[\[FIGURE\s+3\.\d+:\s+(.+)\]\]$", line)
        if figure_match:
            key = figure_match.group(1)
            if key not in FIGURES:
                raise ValueError(f"Unknown figure key: {key}")
            add_figure(document, key, pending_svg_links)
            pending_figure_key = key
            continue

        caption_match = re.match(r"^\*\*Figure\s+(3\.\d+):\*\*\s*(.*)$", line)
        if caption_match:
            if pending_figure_key is None:
                raise ValueError(f"Figure caption without marker: {line}")
            expected = FIGURES[pending_figure_key]["number"]
            if caption_match.group(1) != expected:
                raise ValueError(f"Figure number mismatch: expected {expected}")
            add_caption(document, f"Figure {expected}", caption_match.group(2))
            finish_figure_section(document, pending_figure_key)
            pending_figure_key = None
            continue

        if line.startswith("#### "):
            add_heading(document, line[5:], subsection=True)
            continue
        if line.startswith("### "):
            add_heading(document, line[4:], subsection=False)
            continue
        if in_references:
            add_reference(document, line)
            continue
        if line.startswith("- "):
            add_bullet_item(document, line[2:])
            continue
        if re.match(r"^\d+\.\s", line):
            add_numbered_item(document, line)
            continue
        add_body_paragraph(document, line)

    if pending_figure_key is not None:
        raise ValueError(f"Figure {pending_figure_key} has no caption")


def set_document_properties(document):
    properties = document.core_properties
    properties.title = "Chapter Three — Phishing Digital Medium Detection System Using Machine Learning"
    properties.subject = "Final-Year Project Report — Research Methodology, System Analysis and Design"
    properties.author = "Honour Jesutofaye Jesutofaye"
    properties.keywords = (
        "phishing, smishing, methodology, system analysis, system design, Django, TF-IDF, "
        "Logistic Regression, hybrid detection, explainability, PhishGuard AI"
    )
    properties.comments = (
        "Implementation-grounded Chapter Three. Email dataset provenance and all classifier "
        "performance values remain explicit placeholders pending reproducible evidence."
    )


def set_update_fields(document):
    settings = document.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def inject_svg_fallbacks(docx_path, links):
    """Add each editable SVG as the Office vector source for its PNG fallback."""
    if not links:
        return
    namespaces = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
        "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
        "asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main",
    }
    with tempfile.TemporaryDirectory() as temporary:
        unpacked = Path(temporary) / "docx"
        with zipfile.ZipFile(docx_path) as archive:
            archive.extractall(unpacked)

        document_xml = unpacked / "word" / "document.xml"
        relationships_xml = unpacked / "word" / "_rels" / "document.xml.rels"
        content_types_xml = unpacked / "[Content_Types].xml"
        document_tree = etree.parse(str(document_xml))
        relationships_tree = etree.parse(str(relationships_xml))
        content_types_tree = etree.parse(str(content_types_xml))

        existing_ids = []
        for relationship in relationships_tree.getroot():
            match = re.match(r"rId(\d+)$", relationship.get("Id", ""))
            if match:
                existing_ids.append(int(match.group(1)))
        next_id = max(existing_ids, default=0) + 1

        defaults = content_types_tree.xpath("/ct:Types/ct:Default[@Extension='svg']", namespaces=namespaces)
        if not defaults:
            default = etree.Element(f"{{{namespaces['ct']}}}Default")
            default.set("Extension", "svg")
            default.set("ContentType", "image/svg+xml")
            content_types_tree.getroot().append(default)

        for sequence, (png_rid, svg_source) in enumerate(links, start=1):
            svg_name = f"phishguard-ch3-{sequence}.svg"
            shutil.copyfile(svg_source, unpacked / "word" / "media" / svg_name)
            svg_rid = f"rId{next_id}"
            next_id += 1

            relationship = etree.Element(f"{{{namespaces['pr']}}}Relationship")
            relationship.set("Id", svg_rid)
            relationship.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image")
            relationship.set("Target", f"media/{svg_name}")
            relationships_tree.getroot().append(relationship)

            blips = document_tree.xpath(f"//a:blip[@r:embed='{png_rid}']", namespaces=namespaces)
            if not blips:
                raise ValueError(f"Could not locate fallback image relationship {png_rid}")
            for blip in blips:
                ext_list = etree.SubElement(blip, f"{{{namespaces['a']}}}extLst")
                extension = etree.SubElement(ext_list, f"{{{namespaces['a']}}}ext")
                extension.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
                svg_blip = etree.SubElement(extension, f"{{{namespaces['asvg']}}}svgBlip")
                svg_blip.set(f"{{{namespaces['r']}}}embed", svg_rid)

        document_tree.write(str(document_xml), xml_declaration=True, encoding="UTF-8", standalone="yes")
        relationships_tree.write(str(relationships_xml), xml_declaration=True, encoding="UTF-8", standalone="yes")
        content_types_tree.write(str(content_types_xml), xml_declaration=True, encoding="UTF-8", standalone="yes")

        rebuilt = Path(temporary) / "rebuilt.docx"
        with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(unpacked.rglob("*")):
                if path.is_file():
                    archive.write(path, path.relative_to(unpacked).as_posix())
        os.replace(rebuilt, docx_path)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    for figure in FIGURES.values():
        if not figure["svg"].exists():
            raise FileNotFoundError(figure["svg"])

    document = Document()
    configure_styles(document)
    set_document_properties(document)
    add_cover(document)
    begin_main_section(document)
    pending_svg_links = []
    parse_source(document, SOURCE.read_text(encoding="utf-8").splitlines(), pending_svg_links)
    set_update_fields(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    inject_svg_fallbacks(OUTPUT, pending_svg_links)
    print(OUTPUT)


if __name__ == "__main__":
    main()
