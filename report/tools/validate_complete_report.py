from __future__ import annotations

"""Structural acceptance checks for the combined PhishGuard AI DOCX."""

from collections import Counter
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "report" / "Complete_Project_Report_Phishing_Digital_Medium_Detection_System.docx"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NAMESPACES = {
    "w": W_NS,
    "asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def main():
    require(REPORT.exists(), f"Missing report: {REPORT}")
    with zipfile.ZipFile(REPORT) as archive:
        require(archive.testzip() is None, "The DOCX ZIP package is corrupt")

    document = Document(REPORT)
    nonempty = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    style_counts = Counter(paragraph.style.name for paragraph in document.paragraphs if paragraph.text.strip())

    require(len(document.sections) == 37, f"Expected 37 sections, found {len(document.sections)}")
    require(
        sum(section.orientation == WD_ORIENT.LANDSCAPE for section in document.sections) == 17,
        "Landscape inserts were not preserved",
    )
    require(len(document.tables) == 45, f"Expected 45 total Word tables, found {len(document.tables)}")
    require(style_counts["PCU Chapter Heading"] == 6, "Five chapters plus references heading expected")
    require(style_counts["PCU Chapter Subtitle"] == 5, "Five chapter subtitles expected")
    require(style_counts["PCU Table Caption"] == 32, "Expected 32 numbered table captions")
    require(style_counts["PCU Figure Caption"] == 19, "Expected 19 numbered figure captions")
    require(style_counts["PCU Reference"] == 35, "Expected 35 consolidated APA references")
    require(style_counts["PCU Preliminary Heading"] == 10, "Expected ten preliminary pages")

    for heading in (
        "CERTIFICATION",
        "DECLARATION",
        "APPROVAL",
        "DEDICATION",
        "ACKNOWLEDGEMENTS",
        "ABSTRACT",
        "TABLE OF CONTENTS",
        "LIST OF TABLES",
        "LIST OF FIGURES",
        "LIST OF ABBREVIATIONS",
        "CHAPTER ONE",
        "CHAPTER TWO",
        "CHAPTER THREE",
        "CHAPTER FOUR",
        "CHAPTER FIVE",
        "REFERENCES",
    ):
        require(nonempty.count(heading) == 1, f"Missing or repeated heading: {heading}")

    require(any("[INSERT MATRICULATION NUMBER]" in text for text in nonempty), "Matriculation placeholder is missing")
    require(
        not any("[[TABLE" in text or "[[FIGURE" in text or "**" in text for text in nonempty),
        "Unparsed manuscript markup remains in the DOCX",
    )

    cell_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    require(cell_text.count("EDITABLE SCREENSHOT PLACEHOLDER") == 7, "Expected seven screenshot placeholders")
    require(cell_text.count("EDITABLE RESULT PLACEHOLDER") == 2, "Expected two result placeholders")

    for chapter, count in ((2, 1), (3, 12), (4, 13), (5, 6)):
        for sequence in range(1, count + 1):
            label = f"Table {chapter}.{sequence}:"
            require(sum(text.startswith(label) for text in nonempty) == 1, f"Missing or repeated {label}")
    for chapter, count in ((2, 1), (3, 6), (4, 11), (5, 1)):
        for sequence in range(1, count + 1):
            label = f"Figure {chapter}.{sequence}:"
            require(sum(text.startswith(label) for text in nonempty) == 1, f"Missing or repeated {label}")

    with zipfile.ZipFile(REPORT) as archive:
        document_xml = etree.fromstring(archive.read("word/document.xml"))
        settings_xml = etree.fromstring(archive.read("word/settings.xml"))
        relationships_xml = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
        content_types_xml = etree.fromstring(archive.read("[Content_Types].xml"))

        instructions = document_xml.xpath("//w:instrText/text()", namespaces=NAMESPACES)
        require(sum("TOC" in instruction for instruction in instructions) == 3, "TOC/list fields are missing")
        require(
            settings_xml.xpath('boolean(/w:settings/w:updateFields[@w:val="true"])', namespaces=NAMESPACES),
            "Word is not configured to update fields on open",
        )
        require(len(document_xml.xpath("//asvg:svgBlip", namespaces=NAMESPACES)) == 10, "Expected ten SVG fallbacks")
        require(
            len(
                relationships_xml.xpath(
                    '/pr:Relationships/pr:Relationship[contains(@Target,".svg")]', namespaces=NAMESPACES
                )
            )
            == 10,
            "Expected ten editable SVG relationships",
        )
        require(
            content_types_xml.xpath(
                'boolean(/ct:Types/ct:Default[@Extension="svg" and @ContentType="image/svg+xml"])',
                namespaces=NAMESPACES,
            ),
            "SVG package content type is missing",
        )

        formats = document_xml.xpath("//w:sectPr/w:pgNumType/@w:fmt", namespaces=NAMESPACES)
        starts = document_xml.xpath("//w:sectPr/w:pgNumType/@w:start", namespaces=NAMESPACES)
        alignments = Counter(document_xml.xpath("//w:sectPr/w:vAlign/@w:val", namespaces=NAMESPACES))
        require("lowerRoman" in formats and "decimal" in formats, "Roman/Arabic numbering formats are missing")
        require("2" in starts and "1" in starts, "Preliminary/main page-number starts are missing")
        require(alignments["center"] == 1, "Only the title page should be vertically centred")

        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        require(len([name for name in media if name.endswith(".svg")]) == 10, "Expected ten embedded SVGs")
        require(len([name for name in media if name.endswith(".png")]) == 10, "Expected ten PNG fallbacks")

        footer_instructions = []
        for name in archive.namelist():
            if re.fullmatch(r"word/footer\d+\.xml", name):
                root = etree.fromstring(archive.read(name))
                footer_instructions.extend(root.xpath("//w:instrText/text()", namespaces=NAMESPACES))
        require(sum("PAGE" in instruction for instruction in footer_instructions) == 2, "Expected two PAGE footer fields")

    print(
        f"VALID: {REPORT.name}; {len(nonempty)} nonempty paragraphs; {len(document.tables)} Word tables; "
        f"32 numbered tables; 19 figures; 35 references; {len(document.sections)} sections."
    )


if __name__ == "__main__":
    main()
